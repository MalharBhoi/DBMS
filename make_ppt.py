"""
Generate the DBMS Mini Project presentation.
Run:  source venv/bin/activate && python make_ppt.py
Output: Student_Management_System_PPT.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── Color Palette ──
DARK   = RGBColor(0x2C, 0x3E, 0x50)
ACCENT = RGBColor(0x29, 0x80, 0xB9)
GREEN  = RGBColor(0x27, 0xAE, 0x60)
RED    = RGBColor(0xC0, 0x39, 0x2B)
PURPLE = RGBColor(0x8E, 0x44, 0xAD)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT  = RGBColor(0xEC, 0xF0, 0xF1)
GRAY   = RGBColor(0x7F, 0x8C, 0x8D)


# ══════════════════════════════════════════
#   HELPER FUNCTIONS
# ══════════════════════════════════════════

def add_bg(slide, color=DARK):
    """Fill slide background with a solid color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape_rect(slide, left, top, width, height, color):
    """Add a colored rectangle."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    """Add a text box with specified formatting."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_slide(title_text, bullets, subtitle=""):
    """Create a slide with title and bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide, WHITE)

    # Top accent bar
    add_shape_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.1), DARK)
    add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.8),
                 title_text, font_size=32, color=WHITE, bold=True)

    if subtitle:
        add_text_box(slide, Inches(0.8), Inches(1.3), Inches(11), Inches(0.5),
                     subtitle, font_size=14, color=GRAY, bold=False)

    # Bullets
    y_start = Inches(1.9) if subtitle else Inches(1.5)
    for i, bullet in enumerate(bullets):
        add_text_box(slide, Inches(1.0), y_start + Inches(i * 0.55),
                     Inches(11), Inches(0.5),
                     f"▸  {bullet}", font_size=18, color=DARK)

    # Bottom accent line
    add_shape_rect(slide, Inches(0), Inches(7.3), prs.slide_width, Inches(0.2), ACCENT)
    return slide


# ══════════════════════════════════════════
#   SLIDE 1 — TITLE
# ══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)

add_shape_rect(slide, Inches(0), Inches(2.5), prs.slide_width, Inches(3), ACCENT)

add_text_box(slide, Inches(0), Inches(2.7), prs.slide_width, Inches(1.2),
             "Student Management System", font_size=44, color=WHITE, bold=True,
             alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(0), Inches(3.8), prs.slide_width, Inches(0.6),
             "DBMS Mini Project", font_size=24, color=LIGHT,
             alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(0), Inches(4.5), prs.slide_width, Inches(0.5),
             "Python  |  SQLite  |  Tkinter", font_size=18, color=LIGHT,
             alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(0), Inches(6.2), prs.slide_width, Inches(0.5),
             "Subject: Database Management System (DBMS)", font_size=14,
             color=GRAY, alignment=PP_ALIGN.CENTER)


# ══════════════════════════════════════════
#   SLIDE 2 — INTRODUCTION
# ══════════════════════════════════════════
add_bullet_slide(
    "Introduction",
    [
        "This project is a Student Management System built as a DBMS mini project.",
        "It allows users to Add, View, Search, Update and Delete student records.",
        "Data is stored in a SQLite database (students.db) which persists across sessions.",
        "Two versions are provided: a Terminal-based CLI and a GUI using Tkinter.",
        "The project demonstrates core database concepts — CRUD operations with SQL.",
        "It is designed to be lightweight, with zero external dependencies.",
    ]
)


# ══════════════════════════════════════════
#   SLIDE 3 — TECHNOLOGIES USED
# ══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.1), DARK)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.8),
             "Technologies Used", font_size=32, color=WHITE, bold=True)

techs = [
    ("Python 3",     "Backend programming language",              ACCENT),
    ("SQLite",       "Lightweight embedded relational database",  GREEN),
    ("Tkinter",      "Python's built-in GUI toolkit",             PURPLE),
    ("SQL",          "CREATE, INSERT, SELECT, UPDATE, DELETE",    RED),
]
for i, (name, desc, clr) in enumerate(techs):
    x = Inches(0.8 + i * 3.1)
    y = Inches(2.0)
    box = add_shape_rect(slide, x, y, Inches(2.8), Inches(3.5), clr)
    add_text_box(slide, x, y + Inches(0.5), Inches(2.8), Inches(0.8),
                 name, font_size=26, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), y + Inches(1.5), Inches(2.4), Inches(1.5),
                 desc, font_size=16, color=WHITE, alignment=PP_ALIGN.CENTER)

add_shape_rect(slide, Inches(0), Inches(7.3), prs.slide_width, Inches(0.2), ACCENT)


# ══════════════════════════════════════════
#   SLIDE 4 — FEATURES
# ══════════════════════════════════════════
add_bullet_slide(
    "Project Features",
    [
        "Add New Student — Insert a new student record into the database",
        "View All Students — Display all stored records in a table",
        "Search by Roll Number — Find a specific student using their roll no.",
        "Update Student Details — Modify existing student information",
        "Delete Student Record — Remove a student from the database",
        "Exit Application — Close the program and save all data",
    ],
    subtitle="Complete CRUD (Create, Read, Update, Delete) Operations"
)


# ══════════════════════════════════════════
#   SLIDE 5 — DATABASE DESIGN
# ══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.1), DARK)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.8),
             "Database Design", font_size=32, color=WHITE, bold=True)

add_text_box(slide, Inches(0.8), Inches(1.4), Inches(5), Inches(0.4),
             "Database: students.db  |  Table: student", font_size=14, color=GRAY)

# Table header
headers = ["Column", "Data Type", "Constraint"]
col_widths = [Inches(2.5), Inches(2.5), Inches(3)]
x_start = Inches(1.5)
y_header = Inches(2.0)

for j, (h, w) in enumerate(zip(headers, col_widths)):
    x = x_start + sum(col_widths[:j], Inches(0))
    cell = add_shape_rect(slide, x, y_header, w, Inches(0.5), DARK)
    add_text_box(slide, x, y_header + Inches(0.05), w, Inches(0.45),
                 h, font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Table rows
rows = [
    ("roll",   "INTEGER", "PRIMARY KEY"),
    ("name",   "TEXT",    "NOT NULL"),
    ("course", "TEXT",    "NOT NULL"),
    ("sem",    "INTEGER", "NOT NULL"),
    ("mobile", "TEXT",    "NOT NULL"),
    ("marks",  "REAL",    "NOT NULL"),
]
for i, (col, dtype, constraint) in enumerate(rows):
    y = y_header + Inches(0.5 + i * 0.45)
    bg = LIGHT if i % 2 == 0 else WHITE
    for j, (val, w) in enumerate(zip([col, dtype, constraint], col_widths)):
        x = x_start + sum(col_widths[:j], Inches(0))
        add_shape_rect(slide, x, y, w, Inches(0.45), bg)
        add_text_box(slide, x, y + Inches(0.03), w, Inches(0.4),
                     val, font_size=15, color=DARK, alignment=PP_ALIGN.CENTER,
                     font_name="Consolas")

add_shape_rect(slide, Inches(0), Inches(7.3), prs.slide_width, Inches(0.2), ACCENT)


# ══════════════════════════════════════════
#   SLIDE 6 — SQL QUERIES USED
# ══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.1), DARK)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.8),
             "SQL Queries Used", font_size=32, color=WHITE, bold=True)

queries = [
    ("CREATE TABLE", "CREATE TABLE IF NOT EXISTS student (\n    roll INTEGER PRIMARY KEY, name TEXT, ...)"),
    ("INSERT",       "INSERT INTO student VALUES (?, ?, ?, ?, ?, ?)"),
    ("SELECT ALL",   "SELECT * FROM student"),
    ("SEARCH",       "SELECT * FROM student WHERE roll = ?"),
    ("UPDATE",       "UPDATE student SET name=?, course=? ... WHERE roll=?"),
    ("DELETE",       "DELETE FROM student WHERE roll = ?"),
]
colors = [ACCENT, GREEN, PURPLE, PURPLE, ACCENT, RED]
for i, ((label, query), clr) in enumerate(zip(queries, colors)):
    row = i // 2
    col = i % 2
    x = Inches(0.8 + col * 6.2)
    y = Inches(1.5 + row * 1.8)
    add_shape_rect(slide, x, y, Inches(5.8), Inches(0.4), clr)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.02), Inches(5.4), Inches(0.35),
                 label, font_size=15, color=WHITE, bold=True)
    # code box
    add_shape_rect(slide, x, y + Inches(0.4), Inches(5.8), Inches(1.1), LIGHT)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.5), Inches(5.4), Inches(0.9),
                 query, font_size=13, color=DARK, font_name="Consolas")

add_shape_rect(slide, Inches(0), Inches(7.3), prs.slide_width, Inches(0.2), ACCENT)


# ══════════════════════════════════════════
#   SLIDE 7 — TERMINAL VERSION
# ══════════════════════════════════════════
add_bullet_slide(
    "Version 1: Terminal-Based System",
    [
        "File: student_terminal.py",
        "Runs entirely in the command line / terminal",
        "Provides a numbered menu (1-6) for all operations",
        "User types inputs like Roll Number, Name, etc. when prompted",
        "Output is displayed as a formatted text table",
        "Good for understanding basic Python + SQLite connectivity",
        "Command to run:   python student_terminal.py",
    ],
    subtitle="CLI version — no GUI dependencies needed"
)


# ══════════════════════════════════════════
#   SLIDE 8 — GUI VERSION
# ══════════════════════════════════════════
add_bullet_slide(
    "Version 2: GUI-Based System (Tkinter)",
    [
        "File: student_gui.py",
        "Full graphical window with form fields and buttons",
        "Color-coded buttons: Add (green), Search (purple), Update (blue), Delete (red)",
        "Student records displayed in a scrollable Treeview table",
        "Click any row in the table to auto-fill the form for editing",
        "Status bar at bottom shows total record count",
        "Confirmation dialog before deleting records",
        "Command to run:   python student_gui.py",
    ],
    subtitle="Graphical version using Python's Tkinter library"
)


# ══════════════════════════════════════════
#   SLIDE 9 — HOW IT WORKS (FLOW)
# ══════════════════════════════════════════
add_bullet_slide(
    "How It Works — Application Flow",
    [
        "1.  Application starts → connects to SQLite database (students.db)",
        "2.  If 'student' table doesn't exist → CREATE TABLE is executed",
        "3.  User interacts via menu (terminal) or buttons (GUI)",
        "4.  Each action runs the corresponding SQL query (INSERT / SELECT / UPDATE / DELETE)",
        "5.  Results are displayed in table format (terminal) or Treeview (GUI)",
        "6.  Data is committed to database → persists even after closing the app",
        "7.  On exit → database connection is closed cleanly",
    ],
    subtitle="Step-by-step execution flow"
)


# ══════════════════════════════════════════
#   SLIDE 10 — DOCKER SUPPORT
# ══════════════════════════════════════════
add_bullet_slide(
    "Docker Support",
    [
        "A Dockerfile is included for containerized execution",
        "Uses python:3.11-slim base image (SQLite is built-in, no extra installs)",
        "Build:   docker build -t student-sms .",
        "Run:     docker run -it student-sms",
        "Only the terminal version runs in Docker (GUI needs a display server)",
        "No requirements.txt needed — everything is from Python's standard library",
    ],
    subtitle="Containerized deployment using Docker"
)


# ══════════════════════════════════════════
#   SLIDE 11 — PROJECT STRUCTURE
# ══════════════════════════════════════════
add_bullet_slide(
    "Project File Structure",
    [
        "StudentManagementSystem/",
        "├── student_terminal.py        Terminal-based CRUD application (114 lines)",
        "├── student_gui.py               GUI-based application with Tkinter (180+ lines)",
        "├── Dockerfile                       Docker container for terminal version",
        "├── README.md                     Instructions on how to run the project",
        "└── students.db                    SQLite database (auto-created at runtime)",
    ],
    subtitle="Minimal and clean — only essential files"
)


# ══════════════════════════════════════════
#   SLIDE 12 — CONCLUSION
# ══════════════════════════════════════════
add_bullet_slide(
    "Conclusion",
    [
        "Successfully implemented a complete Student Management System",
        "Demonstrated all CRUD operations using SQL with Python",
        "Built two versions — Terminal and GUI — for versatility",
        "Used SQLite for persistent, file-based data storage",
        "Project requires zero external dependencies (only Python standard library)",
        "Docker support added for easy deployment and portability",
    ]
)


# ══════════════════════════════════════════
#   SLIDE 13 — THANK YOU
# ══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_shape_rect(slide, Inches(0), Inches(2.8), prs.slide_width, Inches(2.2), ACCENT)
add_text_box(slide, Inches(0), Inches(3.0), prs.slide_width, Inches(1),
             "Thank You!", font_size=48, color=WHITE, bold=True,
             alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(0), Inches(4.0), prs.slide_width, Inches(0.6),
             "Questions?", font_size=24, color=LIGHT, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(0), Inches(5.8), prs.slide_width, Inches(0.4),
             "Student Management System — DBMS Mini Project", font_size=14,
             color=GRAY, alignment=PP_ALIGN.CENTER)


# ── Save ──
output = "Student_Management_System_PPT.pptx"
prs.save(output)
print(f"✅ Presentation saved: {output}")
print(f"   Total slides: {len(prs.slides)}")
